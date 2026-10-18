"""Human Office edits and subsequent native work share one canonical version chain."""

import base64
import hashlib
import json
import os
import re
import shlex
import sys
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from sqlalchemy import select

from app.database import pin_rls_tenant_context
from app.models.chat_artifact import ChatArtifact
from app.models.runtime_task import RuntimeTask
from tests.integration.test_m2c_managed_environment_postgres import _exercise_real_worker_environment
from tests.integration.test_m1g24_workspace_artifact_authority_postgres import native_workspace  # noqa: F401
from tests.integration.test_web_chat_worker_restart_kernel_recovery import ScriptedLLMClient

pytestmark = pytest.mark.usefixtures("migrated_pg_url")


async def test_human_office_edits_keep_snapshots_and_feed_the_next_native_work(
    owner_sessionmaker, app_user_sessionmaker, monkeypatch, tmp_path
):
    from app.scripts.verify_officecli_binary_contract import _write_contract_fixtures
    from docx import Document

    inputs = tmp_path / "office-fixtures"
    inputs.mkdir()
    fixtures = _write_contract_fixtures(inputs)
    document = Document(fixtures["docx"])
    document.paragraphs[1].text = "Unchanged paragraph"
    document.save(fixtures["docx"])
    encoded = {f"human.{kind}": base64.b64encode(path.read_bytes()).decode() for kind, path in fixtures.items()}
    create = ("from pathlib import Path; import base64; "
        "Path('saved.txt').write_text('CANONICAL_ENVIRONMENT_FILE'); "
        + f"[Path(name).write_bytes(base64.b64decode(value)) for name,value in {encoded!r}.items()]; "
        + "print('FIRST_COMMAND')")
    command = f"{shlex.quote(str(Path(sys.executable).resolve()))} -I -c {shlex.quote(create)}"
    paths = [f"workspace/human.{kind}" for kind in ("docx", "xlsx", "pptx")]

    async def edit_documents(seed, client, url, scope, current):
        from app.api import office, files, chat_sessions
        from app.runtime import invoker
        from app.services.runtime_task_claim_service import RuntimeTaskClaimService
        from app.services.runtime_task_fence import set_runtime_task_fence, reset_runtime_task_fence
        from app.services.web_chat_runtime import execute_web_chat_run
        from app.services.runtime_terminal_boundary_outbox import RuntimeTerminalBoundaryOutboxService
        from app.services.web_terminal_boundary_processor import WebTerminalBoundaryProcessor

        web = client._transport.app
        for router in (office.router, files.router, chat_sessions.router):
            web.include_router(router, prefix="/api")
        snapshots = {}
        async with app_user_sessionmaker() as db:
            await pin_rls_tenant_context(db, seed['tenant_id'])
            rows = list(await db.scalars(select(ChatArtifact).where(ChatArtifact.runtime_task_id == seed['run_id'],
                ChatArtifact.path.in_(paths))))
            assert len(rows) == 3
            snapshots = {row.id: row.snapshot_hash for row in rows}
        hashes = {}
        for kind, path in zip(("docx", "xlsx", "pptx"), paths):
            endpoint = f"/api/agents/{seed['agent_id']}/office/content"
            response = await client.get(endpoint, params={"path": path})
            assert response.status_code == 200, response.text
            assert response.headers['cache-control'] == 'private, no-store'
            model = response.json()
            field = model['fields'][0]
            changes = [{'node_path': field['node_path'], 'value': f'Human {kind} correction 决定性尾部'}]
            if kind == 'xlsx':
                changes.append({'node_path': model['fields'][1]['node_path'], 'value': '=SUM(17,12)'})
            body = {'expected_current_hash': model['content_hash'], 'changes': changes}
            denied = await client.put(endpoint, params={'path': path, 'operator_view': 'true'}, json=body)
            assert denied.status_code == 403
            saved = await client.put(endpoint, params={'path': path}, json=body)
            assert saved.status_code == 200, saved.text
            hashes[path] = saved.json()['content_hash']
            assert hashes[path] != model['content_hash']
            assert hashlib.sha256((seed['agent_root'] / path).read_bytes()).hexdigest() == hashes[path]
            stale = await client.put(endpoint, params={'path': path}, json=body)
            assert stale.status_code == 409, stale.text
            assert hashlib.sha256((seed['agent_root'] / path).read_bytes()).hexdigest() == hashes[path]
            latest = await client.get(endpoint, params={'path': path})
            assert latest.status_code == 200, latest.text
            assert latest.json()['fields'][0]['value'] == changes[0]['value']
            if kind == 'xlsx':
                assert latest.json()['fields'][1]['value'] == '=SUM(17,12)'
            versions = await client.get(f"/api/agents/{seed['agent_id']}/files/versions", params={'path': path})
            assert versions.status_code == 200, versions.text
            assert len(versions.json()['versions']) >= 2
        from docx import Document
        from pptx import Presentation
        document = Document(seed['agent_root'] / paths[0])
        assert document.paragraphs[0].style.style_id == 'Heading1'
        assert document.paragraphs[1].text == 'Unchanged paragraph'
        slide = Presentation(seed['agent_root'] / paths[2]).slides[0]
        assert slide.shapes[0].width == 8 * 914400
        mounted = os.environ.get('HIVE_M2L_DOM_BRIDGE') == '1'
        if mounted:
            from tests.integration.test_m1g90_company_bot_settings_postgres import _exercise_dom_bridge

            await _exercise_dom_bridge(web, seed, tmp_path, entry_file='e2e/office-content-native.tsx', extra_env={
                'HIVE_OFFICE_SESSION_ID': str(seed['session_id']),
                'HIVE_OFFICE_ARTIFACTS': json.dumps({row.path: str(row.id) for row in rows})})
            match = re.search(r'HIVE_OFFICE_INPUT_ID=([0-9a-f-]+)', (tmp_path / 'dom-native.log').read_text())
            assert match
            input_id = UUID(match.group(1))
        else:
            input_id = uuid4()
            request = {'content': 'Read my saved Office corrections and report them in this work.',
                'input_id': str(input_id), 'idempotency_key': f'session-input:{input_id}', 'permission_mode': 'bypassPermissions'}
            submitted = await client.post(f"/api/agents/{seed['agent_id']}/sessions/{seed['session_id']}/runs", json=request)
            assert submitted.status_code == 201, submitted.text
        async with app_user_sessionmaker() as db:
            await pin_rls_tenant_context(db, seed['tenant_id'])
            rows = list(await db.scalars(select(RuntimeTask).where(RuntimeTask.parent_agent_id == seed['agent_id'],
                RuntimeTask.parent_session_id == str(seed['session_id']), RuntimeTask.id != seed['run_id'],
                RuntimeTask.task_type == 'web_chat_turn')))
            assert len(rows) == 1
            run_id = rows[0].id
            claims = await RuntimeTaskClaimService(db=db, worker_id='office-follow-up', lease_seconds=600).claim_available()
            version = next(row for row in claims if row.id == run_id).claim_version
            await db.commit()
        consume = """
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as E
texts=[]
for kind,member in [('docx','word/document.xml'),('xlsx','xl/worksheets/sheet1.xml'),('pptx','ppt/slides/slide1.xml')]:
 with ZipFile('human.'+kind) as z: text=' '.join(E.fromstring(z.read(member)).itertext())
 assert 'Human '+kind+' correction 决定性尾部' in text, (kind,text)
 texts.append(text)
Path('office-follow-up.md').write_text(chr(10).join(texts))
print('OFFICE_HUMAN_BYTES_CONSUMED')
"""
        follow = f"{shlex.quote(str(Path(sys.executable).resolve()))} -I -c {shlex.quote(consume)}"
        if mounted:
            consume = consume.replace(" texts.append(text)", " assert '_UI_SAVED' in text\n texts.append(text)")
            follow = f"{shlex.quote(str(Path(sys.executable).resolve()))} -I -c {shlex.quote(consume)}"
        script = ScriptedLLMClient([
            {'tool_calls': [{'id': 'office-follow-up-command', 'type': 'function', 'function': {
                'name': 'run_command', 'arguments': json.dumps({'command': follow, 'timeout': 10})}}], 'finish_reason': 'tool_calls'},
            {'content': 'Office corrections consumed.\nDELIVERABLE: workspace/office-follow-up.md', 'finish_reason': 'stop'},
        ])
        monkeypatch.setattr(invoker, 'create_llm_client', lambda **_kwargs: script)
        fence = set_runtime_task_fence(task_id=run_id, worker_id='office-follow-up', claim_version=version)
        try:
            await execute_web_chat_run(run_id)
        finally:
            reset_runtime_task_fence(fence)
        assert 'OFFICE_HUMAN_BYTES_CONSUMED' in json.dumps(script.requests[-1]['messages'])
        processor = WebTerminalBoundaryProcessor(session_factory=app_user_sessionmaker, data_root=seed['data_root'])
        result = await RuntimeTerminalBoundaryOutboxService(session_factory=app_user_sessionmaker).drain_once(
            tenant_id=seed['tenant_id'], worker_id='office-boundary', canonical_validator=processor.validate,
            process_callback=processor, task_types=('web_chat_turn',))
        assert result['delivered'] == 1
        async with app_user_sessionmaker() as db:
            await pin_rls_tenant_context(db, seed['tenant_id'])
            assert (await db.get(RuntimeTask, run_id)).status == 'completed'
            for artifact_id, digest in snapshots.items():
                assert (await db.get(ChatArtifact, artifact_id)).snapshot_hash == digest
            report = await db.scalar(select(ChatArtifact).where(ChatArtifact.runtime_task_id == run_id,
                ChatArtifact.path == 'workspace/office-follow-up.md'))
            assert report is not None
            response = await client.get(f"/api/agents/{seed['agent_id']}/files/artifacts/{report.id}/content")
            assert response.status_code == 200, response.text
            assert 'Human docx correction 决定性尾部' in response.json()['content']
        (tmp_path / 'office-content-observation.json').write_text(json.dumps({'saved_hashes': hashes,
            'original_artifacts': {str(k): v for k, v in snapshots.items()}, 'follow_up_run': str(run_id)}, indent=2))
        return (await client.get(url, params=scope)).json()['environment']

    await _exercise_real_worker_environment(owner_sessionmaker, app_user_sessionmaker, monkeypatch, tmp_path,
        recovery=edit_documents, initial_command=command, initial_artifact_paths=paths,
        recovery_invocation_key='office-follow-up-command')


@pytest.fixture
async def human_document(native_workspace, app_user_sessionmaker):  # noqa: F811
    from io import BytesIO
    from docx import Document
    from app.api import office, files
    from app.models.user import User

    n = native_workspace
    n.client._transport.app.include_router(office.router)
    buffer = BytesIO()
    document = Document()
    document.add_paragraph('Original canonical body')
    document.save(buffer)
    path = 'workspace/race.docx'
    async with app_user_sessionmaker() as db:
        await pin_rls_tenant_context(db, n.tenant)
        user = await db.get(User, n.user)
        result = await files._write_file_bytes(n.agent, path, buffer.getvalue(), expected_current_hash=None,
            current_user=user, db=db)
        assert result['status'] == 'ok'
    return n, path, buffer.getvalue()


@pytest.mark.parametrize('failure', ['command', 'new_schema_issue'])
async def test_office_private_processing_failure_never_changes_canonical_bytes(human_document, monkeypatch, failure):
    from app.services.office_document_service import OfficeDocumentService
    from app.services.officecli_adapter import OfficeCLIExecutionError

    n, path, original = human_document
    endpoint = f'/agents/{n.agent}/office/content'
    model = (await n.client.get(endpoint, params={'path': path})).json()
    actual = OfficeDocumentService.run_apply
    validation = OfficeDocumentService._content_validation_issues
    observed = []

    def fail_after_private_edit(service, relative, **kwargs):
        result = actual(service, relative, **kwargs)
        changed = service.resolve_document_path(relative)
        assert changed.read_bytes() != original
        assert changed.parent != n.root
        observed.append(changed)
        if failure == 'command':
            raise OfficeCLIExecutionError(command='batch', returncode=1, stderr='Synthetic failure after real private edit', payload=result)
        return result

    def new_issue_after_edit(service, relative):
        issues = validation(service, relative)
        if observed:
            issues['Synthetic newly introduced schema issue'] += 1
        return issues

    monkeypatch.setattr(OfficeDocumentService, 'run_apply', fail_after_private_edit)
    if failure == 'new_schema_issue':
        monkeypatch.setattr(OfficeDocumentService, '_content_validation_issues', new_issue_after_edit)
    result = await n.client.put(endpoint, params={'path': path}, json={
        'expected_current_hash': model['content_hash'], 'changes': [
            {'node_path': model['fields'][0]['node_path'], 'value': 'Must remain private until commit'}]})
    assert result.status_code == 503
    assert result.json()['detail']['outcome'] == 'not_applied'
    assert (n.root / path).read_bytes() == original
    assert len(observed) == 1 and not observed[0].exists()


async def test_another_writer_wins_during_private_office_processing(human_document, monkeypatch):
    import asyncio
    import threading
    from app.services.office_document_service import OfficeDocumentService

    n, path, _original = human_document
    endpoint = f'/agents/{n.agent}/office/content'
    model = (await n.client.get(endpoint, params={'path': path})).json()
    actual = OfficeDocumentService.apply_content_changes
    loop = asyncio.get_running_loop()
    staged = asyncio.Event()
    release = threading.Event()
    claimed = False

    def pause_after_private_edit(service, relative, changes):
        nonlocal claimed
        first = not claimed
        claimed = True
        result = actual(service, relative, changes)
        if first:
            loop.call_soon_threadsafe(staged.set)
            assert release.wait(15)
        return result

    monkeypatch.setattr(OfficeDocumentService, 'apply_content_changes', pause_after_private_edit)
    def body(value):
        return {'expected_current_hash': model['content_hash'], 'changes': [
            {'node_path': model['fields'][0]['node_path'], 'value': value}]}
    pending = asyncio.create_task(n.client.put(endpoint, params={'path': path}, json=body('Delayed writer')))
    try:
        await asyncio.wait_for(staged.wait(), 15)
        winner = await n.client.put(endpoint, params={'path': path}, json=body('Concurrent writer'))
        assert winner.status_code == 200, winner.text
        winner_bytes = (n.root / path).read_bytes()
    finally:
        release.set()
        loser = await pending
    assert loser.status_code == 409, loser.text
    assert (n.root / path).read_bytes() == winner_bytes
    current = await n.client.get(endpoint, params={'path': path})
    assert current.json()['fields'][0]['value'] == 'Concurrent writer'


@pytest.mark.parametrize('alias', ['symlink', 'hardlink'])
async def test_office_content_aliases_never_reach_the_document_processor(human_document, monkeypatch, tmp_path, alias):
    import os
    from app.services.office_document_service import OfficeDocumentService

    n, path, original = human_document
    protected = tmp_path / 'other-source.docx'
    protected.write_bytes(original)
    target = n.root / path
    target.unlink()
    if alias == 'symlink':
        target.symlink_to(protected)
    else:
        os.link(protected, target)
    def forbidden(*_args, **_kwargs):
        raise AssertionError('An aliased source cannot enter the Office processor')
    monkeypatch.setattr(OfficeDocumentService, 'editable_content', forbidden)
    endpoint = f'/agents/{n.agent}/office/content'
    read = await n.client.get(endpoint, params={'path': path})
    write = await n.client.put(endpoint, params={'path': path}, json={
        'expected_current_hash': hashlib.sha256(original).hexdigest(),
        'changes': [{'node_path': '/body/p[1]', 'value': 'Forbidden alias edit'}]})
    assert read.status_code == 403, read.text
    assert write.status_code == 403, write.text
    assert protected.read_bytes() == original


@pytest.mark.parametrize('actions,revoke', [([], False), (['read'], False), (['write'], False),
    (['read', 'write'], False), (['read', 'write'], True)])
async def test_human_office_needs_both_effect_and_source_read_authority(
    human_document, owner_sessionmaker, monkeypatch, actions, revoke
):
    import asyncio
    from sqlalchemy import delete
    from app.core.resource_authority import workspace_resource_id
    from app.models.security_audit import ResourcePermission
    from app.models.workspace_resource import WorkspaceResourceManifest
    from app.services.office_document_service import OfficeDocumentService

    n, path, original = human_document
    endpoint = f'/agents/{n.agent}/office/content'
    model = (await n.client.get(endpoint, params={'path': path})).json()
    grant_id = uuid4()
    async with owner_sessionmaker() as db:
        manifest = await db.scalar(select(WorkspaceResourceManifest).where(
            WorkspaceResourceManifest.agent_id == n.agent, WorkspaceResourceManifest.path == path))
        manifest.owner_user_id = n.foreign
        if actions:
            db.add(ResourcePermission(id=grant_id, tenant_id=n.tenant, principal_type='user', principal_id=n.user,
                resource_type='workspace_file', resource_id=workspace_resource_id(n.agent, path), actions=actions, conditions={}))
        await db.commit()
    reader = await n.client.get(endpoint, params={'path': path})
    assert reader.status_code == (200 if 'read' in actions else 403), reader.text
    if reader.status_code == 200:
        assert reader.json()['read_only'] is ('write' not in actions)
    if revoke:
        actual = OfficeDocumentService.apply_content_changes
        loop = asyncio.get_running_loop()
        async def remove_grant():
            async with owner_sessionmaker() as db:
                await db.execute(delete(ResourcePermission).where(ResourcePermission.id == grant_id))
                await db.commit()
        def revoke_after_processing(service, relative, changes):
            result = actual(service, relative, changes)
            asyncio.run_coroutine_threadsafe(remove_grant(), loop).result(timeout=10)
            return result
        monkeypatch.setattr(OfficeDocumentService, 'apply_content_changes', revoke_after_processing)
    saved = await n.client.put(endpoint, params={'path': path}, json={
        'expected_current_hash': model['content_hash'], 'changes': [
            {'node_path': model['fields'][0]['node_path'], 'value': 'Scoped collaborator correction'}]})
    allowed = set(actions) == {'read', 'write'} and not revoke
    assert saved.status_code == (200 if allowed else 403), saved.text
    if not allowed:
        assert (n.root / path).read_bytes() == original
