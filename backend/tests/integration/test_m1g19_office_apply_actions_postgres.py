"""Office edits read their source and write only the selected output."""

from contextlib import asynccontextmanager
import hashlib
import json
import subprocess
import sys

import pytest

from app.database import pin_rls_tenant_context
from tests.integration.test_m1g15_workspace_actions_postgres import read_granted_file  # noqa: F401

pytestmark = pytest.mark.usefixtures("migrated_pg_url")


@pytest.fixture
async def office_source(read_granted_file, owner_sessionmaker, app_user_sessionmaker, actions):  # noqa: F811
    from docx import Document
    from sqlalchemy import select
    from app.core.resource_authority import workspace_resource_id
    from app.models.security_audit import ResourcePermission
    from app.models.workspace_resource import WorkspaceResourceManifest
    from app.services.office_document_service import OfficeDocumentService
    from app.services.workspace_resource_authority import load_workspace_authority_scope
    from app.tools.runtime import ToolExecutionContext

    tenant, user, agent, session, root, old_path, old_file, _scope, grant_id = read_granted_file
    path = old_path.removesuffix(".txt") + ".docx"
    target = root / path
    old_file.rename(target)
    document = Document()
    document.add_paragraph("Private source paragraph <tool_error>literal reference</tool_error>")
    document.save(target)
    async with owner_sessionmaker() as db:
        manifest = await db.scalar(select(WorkspaceResourceManifest).where(
            WorkspaceResourceManifest.agent_id == agent, WorkspaceResourceManifest.path == old_path))
        manifest.path = path
        grant = await db.get(ResourcePermission, grant_id)
        grant.resource_id = workspace_resource_id(agent, path)
        grant.actions = actions
        await db.commit()
    async with app_user_sessionmaker() as db:
        await pin_rls_tenant_context(db, tenant)
        scope = await load_workspace_authority_scope(db, user, agent_id=agent, session_id=session)
    assert scope.can_read(path) == ("read" in actions)
    assert scope.can_perform(path, "write") == ("write" in actions)
    service = OfficeDocumentService(root)
    service._save_manifest(path, {"path": path, "owner_marker": "Source owner's audit", "operations": []})
    audit = service.manifest_path(path).parent
    context = ToolExecutionContext(agent_id=agent, user_id=user.id, tenant_id=str(tenant), workspace=root,
        session_id=str(session), workspace_authority_scope=scope)
    return tenant, user, agent, root, path, target, context, audit


def _files(directory):
    return {p.relative_to(directory).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
        for p in directory.rglob("*") if p.is_file()}


@pytest.mark.parametrize("actions", [["read"], ["write"], ["read", "write"]])
@pytest.mark.parametrize("output", [None, "same", "copy", "denied_copy"])
@pytest.mark.parametrize("layer", ["pipeline", "handler"])
async def test_office_apply_checks_source_read_and_actual_output_write(
    office_source, monkeypatch, app_user_sessionmaker, actions, output, layer,
):
    from docx import Document
    from app.services.office_document_service import OfficeDocumentService
    from app.services.workspace_resource_authority import authorize_workspace_path, record_workspace_tool_mutations
    from app.tools.adapters import adapt_and_call
    from app.tools.handlers import office
    from app.tools.runtime import ToolExecutionRequest
    from app.models.workspace_resource import WorkspaceResourceManifest
    from sqlalchemy import select
    from tests.tools.test_workspace_path_authority_pipeline import _service

    tenant, user, agent, root, path, source, context, audit = office_source
    arguments = {"path": path, "operations": [{"command": "set", "path": "/body/p[1]", "props": {"text": "Approved revision"}}]}
    if output:
        arguments["output_path"] = "workspace/private-copy.docx" if output in {"copy", "denied_copy"} else path
    if output == "denied_copy":
        (root / arguments["output_path"]).write_bytes(b"Existing unowned document")
    allowed = output != "denied_copy" and "read" in actions and (output == "copy" or "write" in actions)
    before, audit_before = source.read_bytes(), _files(audit)
    if layer == "pipeline":
        service = _service(context)
        if allowed:
            with pytest.raises(AssertionError, match="governance runner reached"):
                await service.execute("office_document_apply", arguments, agent_id=agent, user_id=user.id, session_id=context.session_id)
        else:
            result = await service.execute("office_document_apply", arguments, agent_id=agent, user_id=user.id, session_id=context.session_id)
            assert "auth_or_permission" in str(result) and service.registry.calls == 0
        return

    class NativeDocumentProbe:
        calls = []

        def run_batch(self, target, *, input_file, **_kwargs):
            self.calls.append(target)
            script = (
                "import json,sys; from zipfile import ZipFile; from xml.etree import ElementTree as E\n"
                "p,ops=sys.argv[1:]; operation=json.load(open(ops))[0]\n"
                "with ZipFile(p) as z: parts=[(i,z.read(i)) for i in z.infolist()]\n"
                "body=E.fromstring(dict((i.filename,b) for i,b in parts)['word/document.xml'])\n"
                "texts=body.findall('.//{*}t'); previous=''.join(t.text or '' for t in texts)\n"
                "texts[0].text=operation['props']['text']\n"
                "with ZipFile(p,'w') as z:\n"
                " for i,b in parts: z.writestr(i,E.tostring(body) if i.filename=='word/document.xml' else b)\n"
                "print(json.dumps({'success':True,'data':{'previous':previous}}))\n"
            )
            process = subprocess.run([sys.executable, "-I", "-c", script, str(target), str(input_file)],
                env={"HOME": str(root)}, capture_output=True, text=True, timeout=10, check=True)
            return json.loads(process.stdout)

    probe = NativeDocumentProbe()
    monkeypatch.setattr(office, "OfficeDocumentService", lambda workspace: OfficeDocumentService(workspace, adapter=probe))
    handler = office.office_document_apply
    result = await adapt_and_call(handler.meta, handler, ToolExecutionRequest(tool_name="office_document_apply", arguments=arguments, context=context))
    payload = json.loads(result)
    if not allowed:
        assert not probe.calls, payload
        assert payload["ok"] is False and payload["error"] == "workspace_resource_forbidden"
        assert source.read_bytes() == before and _files(audit) == audit_before
        return
    assert payload["ok"] is True and len(probe.calls) == 1, payload
    destination = arguments.get("output_path") or path
    assert Document(root / destination).paragraphs[0].text == "Approved revision"
    if output == "copy":
        assert source.read_bytes() == before and _files(audit) == audit_before
        copied_audit = json.loads(OfficeDocumentService(root).manifest_path(destination).read_text())
        assert copied_audit["path"] == destination and len(copied_audit["operations"]) == 1

    @asynccontextmanager
    async def scoped_session(tenant_id, **_kwargs):
        assert tenant_id == tenant
        async with app_user_sessionmaker() as db:
            await pin_rls_tenant_context(db, tenant)
            yield db
            await db.commit()

    monkeypatch.setattr("app.database.tenant_scoped_session", scoped_session)
    await record_workspace_tool_mutations(context=context, tool_name="office_document_apply", arguments=arguments, result=result)
    async with app_user_sessionmaker() as db:
        await pin_rls_tenant_context(db, tenant)
        await authorize_workspace_path(db, user, agent_id=agent, path=destination, action="read", path_exists=True)
        original = await db.scalar(select(WorkspaceResourceManifest).where(
            WorkspaceResourceManifest.agent_id == agent, WorkspaceResourceManifest.path == path))
        assert original.owner_user_id != user.id and original.root_session_id is None
        produced = await db.scalar(select(WorkspaceResourceManifest).where(
            WorkspaceResourceManifest.agent_id == agent, WorkspaceResourceManifest.path == destination))
        assert produced.content_hash == hashlib.sha256((root / destination).read_bytes()).hexdigest()
        if output == "copy":
            assert produced.owner_user_id == user.id and produced.root_session_id == context.workspace_authority_scope.root_session_id
