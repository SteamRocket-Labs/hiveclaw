"""File upload API for chat — saves files to agent workspace and extracts text."""

import base64
import hashlib
import os
import subprocess
import uuid
from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, Depends, File, HTTPException, UploadFile, Form
from fastapi.responses import FileResponse
from starlette.background import BackgroundTask
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.permissions import check_agent_access
from app.config import get_settings
from app.core.security import get_current_user
from app.database import get_db, tenant_scoped_session
from app.models.user import User
from app.services.personal_knowledge_service import PersonalKnowledgeService

router = APIRouter(prefix="/chat", tags=["chat"])


class _OriginalDownloadResponse(FileResponse):
    def __init__(self, lease, path, filename):
        self.lease = lease
        super().__init__(path, filename=filename, media_type="application/octet-stream", background=BackgroundTask(lease.cleanup))

    async def __call__(self, scope, receive, send):
        try:
            await super().__call__(scope, receive, send)
        finally:
            self.lease.cleanup()


@router.get("/attachments/{ref_id}/download")
async def download_original_attachment(ref_id: uuid.UUID, current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db)):
    from app.services.blob_store import BlobStoreError
    from app.services.upload_originals import capture_original_download

    try:
        lease, target, filename = await capture_original_download(db, current_user, ref_id)
    except (BlobStoreError, OSError):
        raise HTTPException(503, detail={"code": "original_attachment_unavailable", "retryable": True})
    return _OriginalDownloadResponse(lease, target, filename)

_settings = get_settings()
WORKSPACE_ROOT = Path(_settings.AGENT_DATA_DIR)
CHAT_UPLOAD_MAX_BYTES = int(os.getenv("HIVE_CHAT_UPLOAD_MAX_BYTES", str(25 * 1024 * 1024)))
CHAT_IMAGE_UPLOAD_MAX_BYTES = min(
    CHAT_UPLOAD_MAX_BYTES,
    int(os.getenv("HIVE_CHAT_IMAGE_UPLOAD_MAX_BYTES", str(10 * 1024 * 1024))),
)

# Supported extensions and their text extraction method
TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".xml",
    ".yaml",
    ".yml",
    ".py",
    ".js",
    ".ts",
    ".html",
    ".css",
    ".sql",
    ".sh",
    ".log",
    ".ini",
    ".cfg",
    ".conf",
    ".env",
    ".toml",
}
OFFICE_EXTENSIONS = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".pptx", ".ppt"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
EXTRACTABLE = TEXT_EXTENSIONS | OFFICE_EXTENSIONS

MIME_MAP = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
}


def extract_text(file_path: Path, extension: str) -> str:
    """Extract text content from a file."""
    if extension in TEXT_EXTENSIONS:
        try:
            return file_path.read_text(encoding="utf-8", errors="replace")
        except Exception:
            return file_path.read_text(encoding="gbk", errors="replace")

    if extension == ".pdf":
        try:
            import pdfplumber

            with pdfplumber.open(file_path) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            return text or "[PDF内容提取失败]"
        except ImportError:
            try:
                result = subprocess.run(
                    ["pdftotext", str(file_path), "-"],
                    capture_output=True,
                    text=True,
                    timeout=30,
                )
                return result.stdout.strip() or "[无法解析PDF]"
            except Exception as e:
                return f"[PDF解析错误: {e}]"
        except Exception as e:
            return f"[PDF解析错误: {e}]"

    if extension == ".docx":
        try:
            from docx import Document

            doc = Document(file_path)
            text = "\n".join(p.text for p in doc.paragraphs)
            return text or "[DOCX内容提取失败]"
        except ImportError:
            return "[需要安装 python-docx 库]"
        except Exception as e:
            return f"[DOCX解析错误: {e}]"

    if extension in (".xlsx", ".xls"):
        try:
            import openpyxl

            wb = openpyxl.load_workbook(file_path, read_only=True)
            lines: list[str] = []
            for ws in wb.worksheets:
                lines.append(f"## Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    lines.append("\t".join(str(cell) if cell is not None else "" for cell in row))
            return "\n".join(lines) or "[Excel内容提取失败]"
        except ImportError:
            return "[需要安装 openpyxl 库]"
        except Exception as e:
            return f"[Excel解析错误: {e}]"

    return f"[不支持的文件格式: {extension}]"


def _sanitize_upload_filename(filename: str) -> str:
    """Normalize a user-supplied filename to a safe basename."""
    normalized = filename.replace("\\", "/")
    safe_name = Path(normalized).name
    if safe_name in {"", ".", ".."} or "\x00" in safe_name:
        raise HTTPException(status_code=400, detail="Invalid filename")
    return safe_name


def _write_upload_copy(namespace, filename, content):
    from app.services.confined_storage import confined_storage_parent

    counter = 0
    original = Path(filename)
    while True:
        name = filename if counter == 0 else f"{original.stem}_{counter}{original.suffix}"
        relative = f"workspace/uploads/{name}"
        storage_path = f"{namespace}/{relative}"
        try:
            with confined_storage_parent(WORKSPACE_ROOT, storage_path, create=True) as (directory, leaf):
                descriptor = os.open(leaf, os.O_WRONLY | os.O_CREAT | os.O_EXCL | os.O_NOFOLLOW | os.O_CLOEXEC,
                    0o600, dir_fd=directory)
                with os.fdopen(descriptor, "wb") as output:
                    output.write(content)
                    output.flush()
                    os.fsync(output.fileno())
            return relative, WORKSPACE_ROOT / storage_path
        except FileExistsError:
            counter += 1
        except OSError:
            raise HTTPException(503, detail={"code": "upload_workspace_unavailable", "retryable": True})


@router.post("/upload")
async def upload_file(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    agent_id: uuid.UUID | None = Form(None),
    skip_personal_kb: bool = Form(False),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Upload a file for chat context. Saves to agent workspace/uploads/ and returns extracted text."""
    agent = None
    if agent_id:
        access_result = await check_agent_access(db, current_user, agent_id)
        if isinstance(access_result, tuple) and access_result:
            agent = access_result[0]
        elif access_result is not None:
            agent = access_result
    owner_user_id = _owner_user_id_for_upload(agent=agent, current_user=current_user)
    result = await save_upload_for_agent(
        file=file,
        agent_id=agent_id,
        tenant_id=getattr(current_user, "tenant_id", None),
        user_id=getattr(current_user, "id", None),
        db=db,
        personal_kb_enabled=not skip_personal_kb,
        owner_user_id=owner_user_id,
        skip_personal_kb=skip_personal_kb,
        background_tasks=background_tasks,
    )
    if isinstance(db, AsyncSession) and db.in_transaction():
        await db.commit()
    return result


def _owner_user_id_for_upload(*, agent, current_user: User) -> uuid.UUID | None:
    if agent is not None:
        for attr in ("owner_user_id", "creator_id"):
            value = getattr(agent, attr, None)
            if value is not None:
                return uuid.UUID(str(value))
    user_id = getattr(current_user, "id", None)
    return uuid.UUID(str(user_id)) if user_id is not None else None


def _personal_kb_candidate_payload(result, *, origin: str) -> dict[str, object]:
    return {
        "skipped": False,
        "document_id": str(result.document_id),
        "job_id": str(result.job_id) if result.job_id else None,
        "status": str(result.status),
        "warnings": list(result.warnings or []),
        "origin": origin,
    }


async def _process_upload_personal_import_jobs(*, tenant_id: uuid.UUID, owner_user_id: uuid.UUID) -> None:
    # Two-phase worker: the claim commits before long conversion starts.
    await PersonalKnowledgeService().process_import_jobs(
        None,
        session_factory=lambda: tenant_scoped_session(tenant_id),
        tenant_id=tenant_id,
        owner_user_id=owner_user_id,
        current_user_id=owner_user_id,
        limit=5,
        statuses=("queued",),
    )


async def save_upload_for_agent(
    *,
    file: UploadFile,
    agent_id: uuid.UUID | None,
    tenant_id: uuid.UUID | None,
    user_id: uuid.UUID | None,
    db: AsyncSession | None = None,
    personal_kb_enabled: bool = False,
    owner_user_id: uuid.UUID | None = None,
    skip_personal_kb: bool = False,
    background_tasks: BackgroundTasks | None = None,
):
    """Save a chat upload for an already-authorized agent/user context."""
    if not file.filename:
        raise HTTPException(status_code=400, detail="No filename")

    safe_filename = _sanitize_upload_filename(file.filename)
    ext = os.path.splitext(safe_filename)[1].lower()

    content = await file.read()
    if len(content) > CHAT_UPLOAD_MAX_BYTES:
        raise HTTPException(status_code=413, detail=f"File too large (max {CHAT_UPLOAD_MAX_BYTES} bytes)")
    if ext in IMAGE_EXTENSIONS and len(content) > CHAT_IMAGE_UPLOAD_MAX_BYTES:
        raise HTTPException(status_code=413, detail=f"Image too large (max {CHAT_IMAGE_UPLOAD_MAX_BYTES} bytes)")

    original = None
    if isinstance(db, AsyncSession) and tenant_id is not None and user_id is not None:
        from app.services.blob_store import BlobStoreError
        from app.services.upload_originals import archive_upload

        try:
            original = await archive_upload(db, tenant_id=uuid.UUID(str(tenant_id)), user_id=uuid.UUID(str(user_id)),
                agent_id=uuid.UUID(str(agent_id)) if agent_id else None, filename=safe_filename, body=content)
        except (BlobStoreError, OSError):
            raise HTTPException(503, detail={"code": "original_attachment_storage_unavailable", "retryable": True})

    if agent_id:
        from app.services.session_workspace_snapshot import async_agent_workspace_lock

        async with async_agent_workspace_lock(agent_id):
            workspace_path, save_path = _write_upload_copy(str(agent_id), safe_filename, content)
    else:
        tenant_segment = str(tenant_id) if tenant_id else "no_tenant"
        user_segment = str(user_id) if user_id else "anonymous"
        namespace = f"local_agents/{tenant_segment}/users/{user_segment}"
        workspace_path, save_path = _write_upload_copy(namespace, safe_filename, content)
    uploads_dir = save_path.parent

    if agent_id is not None and db is not None and tenant_id is not None and user_id is not None:
        from app.services.workspace_resource_authority import register_workspace_path

        # Receipt failure quarantines the copy. Its path may have changed while
        # awaiting the database, so it cannot safely authorize an unlink.
        await register_workspace_path(
            db,
            tenant_id=uuid.UUID(str(tenant_id)),
            agent_id=uuid.UUID(str(agent_id)),
            path=workspace_path,
            owner_user_id=uuid.UUID(str(user_id)),
            root_session_id=None,
            source="chat_upload",
            content_hash=hashlib.sha256(content).hexdigest(),
            allow_owner_rebind=False,
        )
        if isinstance(db, AsyncSession):
            from app.services.content_input_receipts import (
                latest_workspace_receipt, lock_workspace_receipt_write, record_workspace_receipt,
                validate_browser_input_dependencies,
            )

            await lock_workspace_receipt_write(db, tenant_id, agent_id, workspace_path)
            previous = await latest_workspace_receipt(db, tenant_id, agent_id, workspace_path)
            if previous is not None:
                reader = await db.get(User, user_id)
                await validate_browser_input_dependencies(db, tenant_id, reader, previous.dependencies_json)
            record_workspace_receipt(db, tenant_id, agent_id, workspace_path,
                hashlib.sha256(content).hexdigest(), previous)
            await db.flush()

    # Convert text/documents to canonical Markdown artifacts. Keep image handling
    # separate because vision-capable chat models still consume typed image parts.
    is_image = ext in IMAGE_EXTENSIONS
    image_data_url = ""
    conversion = None
    preview_text = ""
    if is_image:
        # For images: generate base64 data URL for vision models
        mime = MIME_MAP.get(ext, "image/png")
        b64 = base64.b64encode(content).decode("ascii")
        image_data_url = f"data:{mime};base64,{b64}"
        extracted = f"[图片文件: {file.filename}，需要视觉模型分析]"
    elif ext in EXTRACTABLE:
        try:
            from app.services.document_conversion import (
                DocumentConversionRequest,
                DocumentConversionService,
                render_conversion_preview,
            )

            workspace_root = (
                (WORKSPACE_ROOT / str(agent_id)).resolve() if agent_id else uploads_dir.parent.parent.resolve()
            )
            converted = DocumentConversionService().convert(
                DocumentConversionRequest(
                    source_path=save_path,
                    workspace_root=workspace_root,
                    source_uri=None,
                    tenant_id=tenant_id,
                    agent_id=agent_id,
                    user_id=user_id,
                    mode="auto",
                    force_refresh=False,
                )
            )
            extracted = render_conversion_preview(converted, max_chars=None)
            preview_text = render_conversion_preview(converted, max_chars=6000)
            conversion = {
                "status": "converted",
                "markdown_path": converted.artifact_markdown_path,
                "metadata_path": converted.artifact_metadata_path,
                "engine": converted.engine,
                "used_ocr": converted.used_ocr,
                "used_vision": converted.used_vision,
                "warnings": list(converted.warnings),
                "source_sha256": converted.source_sha256,
            }
        except Exception as e:
            extracted = f"[文档转换失败: {e}]"
            preview_text = extracted
            conversion = {
                "status": "failed",
                "error": str(e),
                "warnings": ["document_conversion_failed"],
            }
    else:
        extracted = f"[文件已保存，格式 {ext} 暂不支持文本提取，Agent 可通过 read_document 工具读取]"

    personal_kb_candidate: dict[str, object] = (
        {"skipped": True, "reason": "user_skip"}
        if skip_personal_kb
        else {
            "skipped": True,
            "reason": "not_requested",
        }
    )
    if personal_kb_enabled and db is not None and tenant_id is not None and owner_user_id is not None:
        origin = f"agent:{agent_id}" if agent_id else f"user:{user_id}"
        try:
            kb_result = await PersonalKnowledgeService().queue_source_bytes_import(
                db,
                tenant_id=uuid.UUID(str(tenant_id)),
                owner_user_id=uuid.UUID(str(owner_user_id)),
                filename=safe_filename,
                data=content,
                title=Path(safe_filename).stem or safe_filename,
                source_kind="chat_attachment",
                source_uri=f"{origin}:{workspace_path}",
                source_mime_type=file.content_type,
                created_by_user_id=uuid.UUID(str(user_id)) if user_id is not None else None,
                agent_searchable=True,
                sensitivity="internal",
                doc_metadata={
                    "origin": origin,
                    "workspace_path": workspace_path,
                    "source_context": "chat_upload",
                    "t0_refs": [],
                },
            )
            personal_kb_candidate = _personal_kb_candidate_payload(kb_result, origin=origin)
            if hasattr(db, "commit"):
                await db.commit()
            if background_tasks is not None:
                background_tasks.add_task(
                    _process_upload_personal_import_jobs,
                    tenant_id=uuid.UUID(str(tenant_id)),
                    owner_user_id=uuid.UUID(str(owner_user_id)),
                )
        except Exception as exc:
            personal_kb_candidate = {
                "skipped": False,
                "status": "failed",
                "origin": origin,
                "error": str(exc),
            }

    from app.services.upload_originals import original_upload_projection

    return {
        "filename": safe_filename,
        "saved_filename": save_path.name,
        "size": len(content),
        "extracted_text": extracted,
        "preview_text": preview_text or extracted,
        "conversion": conversion,
        "workspace_path": workspace_path,
        "is_image": is_image,
        "image_data_url": image_data_url,
        "personal_kb_candidate": personal_kb_candidate,
        "original_attachment": original_upload_projection(original) if original is not None else None,
    }
